"""Reservas del día desde Lodgify -> Excel + PDF. Uso: LODGIFY_API_KEY=... python3 reservas_diarias.py [YYYY-MM-DD]"""
import os, sys, json, urllib.request, urllib.parse
from datetime import date, datetime, timedelta, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

KEY = os.environ["LODGIFY_API_KEY"]
BASE = "https://api.lodgify.com/v2"
OUT = os.path.join(os.path.dirname(__file__), "..", "reportes")
PLAT = {"AirbnbIntegration": "Airbnb", "BookingCom": "Booking.com", "Manual": "Directa", "Vrbo": "Vrbo"}

def get(path, **q):
    req = urllib.request.Request(f"{BASE}{path}?{urllib.parse.urlencode(q)}",
                                 headers={"X-ApiKey": KEY, "accept": "application/json"})
    return json.load(urllib.request.urlopen(req, timeout=60))

def main():
    hoy = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else (datetime.now(timezone.utc) - timedelta(hours=6)).date()
    props = {p["id"]: p["name"] for p in get("/properties", size=100)["items"]}
    rows, page = [], 1
    while True:
        d = get("/reservations/bookings", size=50, page=page, includeCount="false",
                updatedSince=(hoy - timedelta(days=3)).isoformat())
        items = d["items"]
        for b in items:
            if b["created_at"][:10] != hoy.isoformat() or b["status"] in ("Declined", "Open") or b.get("is_deleted"):
                continue
            nombre = props.get(b["property_id"], str(b["property_id"]))
            tipo = "Condo" if "condo" in nombre.lower() else "Villa"
            n = (date.fromisoformat(b["departure"]) - date.fromisoformat(b["arrival"])).days
            total = b["total_amount"] or round(sum(v for v in b["subtotals"].values() if v), 2)
            src = b["source"]
            rows.append([b["arrival"], f"{tipo} - {nombre.split('|')[0].strip()}", b["guest"]["name"],
                         PLAT.get(src, src), n, total, b["currency_code"]])
        if len(items) < 50:
            break
        page += 1
    rows.sort(key=lambda r: r[0])
    hdr = ["Fecha de ingreso", "Villa o Condo", "Nombre", "Plataforma", "Noches", "Monto a pagar", "Moneda"]

    os.makedirs(OUT, exist_ok=True)
    xlsx = os.path.join(OUT, f"Reservas_{hoy}.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Reservas del día"
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows: ws.append(r)
    for col, w in zip("ABCDEFG", (16, 40, 28, 14, 8, 16, 8)): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].number_format = '#,##0.00'
    wb.save(xlsx)

    pdf = os.path.join(OUT, f"Ingresos_{hoy}.pdf")
    st = getSampleStyleSheet()
    el = [Paragraph(f"Reporte de ingresos diarios – {hoy:%d/%m/%Y}", st["Title"])]
    total = sum(r[5] for r in rows)
    el.append(Paragraph(f"Reservas: {len(rows)} · Noches: {sum(r[4] for r in rows)} · Total: USD {total:,.2f}", st["Normal"]))
    el.append(Spacer(1, 12))
    por = {}
    for r in rows: por[r[3]] = por.get(r[3], 0) + r[5]
    data = [hdr[:6]] + [r[:6] for r in rows] if rows else [["Sin reservas nuevas hoy"]]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), .3, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    el.append(t); el.append(Spacer(1, 12))
    for p, v in por.items(): el.append(Paragraph(f"{p}: USD {v:,.2f}", st["Normal"]))
    SimpleDocTemplate(pdf, pagesize=landscape(letter)).build(el)
    print(xlsx, pdf, len(rows), total)

main()
