"""Pruebas con el reporte CSV real de Lodgify: python -m pytest automatizacion/tests"""
import sys
from datetime import date
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "automatizacion"))
import lodgify_sync as ls  # noqa: E402

CSV = RAIZ / "OTCReport_512147 (5).csv"
LIBRO = RAIZ / "Canavida_Gastos_Ingresos.xlsx"


def correr(tmp_path, hoy, entrada=LIBRO):
    salida = tmp_path / f"libro_{hoy}.xlsx"
    assert ls.main(["--csv", str(CSV), "--excel", str(entrada), "--salida", str(salida),
                    "--hoy", hoy, "--sin-correo"]) == 0
    return salida


def filas_lodgify(ruta):
    ws = openpyxl.load_workbook(ruta)["Ingresos Diarios"]
    return [ws.cell(f, 10).value for f in range(6, 2027)
            if "Lodgify #" in str(ws.cell(f, 10).value or "")]


def test_montos_como_en_el_libro():
    r = {x.id: x for x in ls.leer_csv(CSV)}
    assert r["21985531"].monto == 319.78        # Booking: total con impuesto
    assert r["21920356"].monto == 197.06        # Expedia: sin impuesto
    assert r["21985531"].villa == "Iguana (Condo 2)"
    assert r["23063777"].villa == "Love Nest"


def test_no_duplica_y_respeta_orden(tmp_path):
    primera = correr(tmp_path, "2026-10-29")
    segunda = correr(tmp_path, "2026-10-29", entrada=primera)
    notas = filas_lodgify(segunda)
    assert notas == filas_lodgify(primera)
    assert len(notas) == len(set(notas)) >= 15  # sin duplicados
    ws = openpyxl.load_workbook(segunda)["Ingresos Diarios"]
    fechas = [ls.parse_fecha(ws.cell(f, 1).value) for f in range(27, ls.ultima_fila_con_datos(ws) + 1)]
    assert fechas == sorted(fechas)


def test_futuras_entran_al_llegar(tmp_path):
    primera = correr(tmp_path, "2026-10-29")
    assert all(ls.parse_fecha(d) <= date(2026, 10, 29) for d in
               [c.value for c in openpyxl.load_workbook(primera)["Ingresos Diarios"]["A"][26:200] if c.value])
    despues = correr(tmp_path, "2026-10-31", entrada=primera)
    assert len(filas_lodgify(despues)) > len(filas_lodgify(primera))


def test_sin_errores_ref(tmp_path):
    wb = openpyxl.load_workbook(correr(tmp_path, "2026-10-29"))
    for ws in wb:
        for fila in ws.iter_rows():
            for c in fila:
                assert "#REF!" not in str(c.value or ""), f"{ws.title}!{c.coordinate}"
