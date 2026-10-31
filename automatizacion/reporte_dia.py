"""Reporte de reservas de un día (por fecha de llegada) desde el CSV de Lodgify -> Excel + PDF."""
import csv, sys
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

csv_path, dia, out = sys.argv[1], sys.argv[2], Path(sys.argv[3])  # dia = YYYY-MM-DD
res = {}
for x in csv.DictReader(open(csv_path, encoding="utf-8-sig")):
    r = res.setdefault(x["Id"], dict(x, total=0.0))
    r["total"] += float(x["GrossAmount"] or 0)
filas = []
for r in res.values():
    if r["Status"] != "Booked":
        continue
    ll = datetime.strptime(r["DateArrival"], "%m/%d/%Y").date()
    if ll.isoformat() != dia:
        continue
    filas.append((ll, r["HouseName"].split("|")[0].strip(), r["GuestName"], r["Source"], int(r["Nights"]), r["total"]))
sal = [(datetime.strptime(r["DateDeparture"], "%m/%d/%Y").date(), r["HouseName"].split("|")[0].strip(), r["GuestName"], r["Source"], int(r["Nights"]), r["total"])
       for r in res.values() if r["Status"] == "Booked" and datetime.strptime(r["DateDeparture"], "%m/%d/%Y").date().isoformat() == dia]
sal.sort(key=lambda t: (t[1], t[2]))
filas.sort(key=lambda t: (t[0], t[1], t[2]))
enc = ["Fecha de ingreso", "Villa / Condo", "Nombre", "Plataforma", "Noches", "Monto a pagar (USD)"]
total = sum(f[5] for f in filas)
out.mkdir(exist_ok=True)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Reservas del día"
ws.append(enc)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
for f in filas:
    ws.append([f[0].strftime("%d/%m/%Y"), *f[1:]])
ws.append(["", "", "", "", "TOTAL", total])
for c in ws[ws.max_row]: c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
    for c in row: c.number_format = '"$"#,##0.00'
for col, w in zip("ABCDEF", (16, 34, 28, 14, 9, 20)): ws.column_dimensions[col].width = w
wb.save(out / f"reservas_{dia}.xlsx")
st = getSampleStyleSheet()
d = datetime.fromisoformat(dia)
el = [Paragraph(f"Canavida · Ingresos diarios — {d:%d/%m/%Y}", st["Title"]),
      Paragraph(f"Reservas con llegada este día: {len(filas)} · Total: ${total:,.2f}", st["Normal"]), Spacer(1, 12)]
data = [enc] + [[f[0].strftime("%d/%m/%Y"), f[1], f[2], f[3], f[4], f"${f[5]:,.2f}"] for f in filas] + [["", "", "", "", "TOTAL", f"${total:,.2f}"]]
t = Table(data, repeatRows=1)
t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
    ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), .3, colors.grey), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ("ALIGN", (4, 0), (-1, -1), "RIGHT")]))
el.append(t)
if not filas:
    el.insert(2, Paragraph("No hay entradas (check-in) registradas para este día.", st["Normal"]))
el += [Spacer(1, 16), Paragraph(f"Salidas (check-out) del día: {len(sal)}", st["Heading3"])]
t2 = Table([["Salida", "Villa / Condo", "Huésped", "Plataforma", "Noches", "Monto"]] + [[f[0].strftime("%d/%m/%Y"), f[1], f[2], f[3], f[4], f"${f[5]:,.2f}"] for f in sal], repeatRows=1)
t2.setStyle(t._tblStyle if hasattr(t, "_tblStyle") else TableStyle([]))
t2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7F7F7F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
    ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), .3, colors.grey), ("ALIGN", (4, 0), (-1, -1), "RIGHT")]))
el.append(t2)
SimpleDocTemplate(str(out / f"ingresos_diarios_{dia}.pdf"), pagesize=letter).build(el)
for f in filas: print(f)
print(total)
